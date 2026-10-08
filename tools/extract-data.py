#!/usr/bin/env python3
"""Build data/ntb-data.json from the "Kalkulator NTB" workbook.

Usage: python3 tools/extract-data.py "/path/to/Kalkulator NTB_SharedV1.1.xlsx"
Reads the KBLI master and the thresholds of Sheet2. No respondent data is read.
"""
import json
import sys
import warnings

import openpyxl

warnings.filterwarnings('ignore')
path = sys.argv[1]
with open(path, 'rb') as handle:  # also works when the file has no .xlsx extension
    wb = openpyxl.load_workbook(handle, data_only=True)
master, s2 = wb['MASTER KBLI'], wb['Sheet2']


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


categories, kbli = [], []
for r in range(2, master.max_row + 1):
    cat, digit, code, title = (master.cell(r, c).value for c in range(1, 5))
    if digit == 1:
        categories.append({'code': code, 'title': title})
    elif digit == 5:
        kbli.append({'c': str(code), 't': title, 'k': cat})

# Sheet2 A:C = NTB ratio range per category ("-" means no range), V:Y = wage percent per category
ntb_range, wage_percent = {}, {}
for r in range(2, 23):
    cat = s2.cell(r, 1).value
    lo, hi = num(s2.cell(r, 2).value), num(s2.cell(r, 3).value)
    ntb_range[cat] = [lo, hi] if lo is not None and hi is not None else None
    wcat = s2.cell(r, 22).value
    if wcat:
        wage_percent[wcat] = [num(s2.cell(r, c).value) for c in (23, 24, 25)]

data = {
    'categories': categories,
    'kbli': kbli,
    'ntbRange': ntb_range,
    'wagePercent': wage_percent,  # [min, mid, max] percent of nilai tambah
    'sizeClasses': [
        {'label': s2.cell(r, 5).value, 'min': int(s2.cell(r, 6).value), 'max': int(s2.cell(r, 7).value) if s2.cell(r, 7).value else None}
        for r in (2, 3, 4)
    ],
    'wageToNtb': [num(s2['J2'].value), num(s2['K2'].value)],
    'wagePerWorkerYear': [num(s2['N2'].value), num(s2['O2'].value)],
    'costShare': {
        'produksi': [num(s2['AB3'].value), num(s2['AC3'].value)],
        'pembelian': [num(s2['AD3'].value), num(s2['AE3'].value)],
        'operasional': [num(s2['AF3'].value), num(s2['AG3'].value)],
        'nonOperasional': [num(s2['AH3'].value), num(s2['AI3'].value)],
    },
}
with open('data/ntb-data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
print(len(categories), 'categories,', len(kbli), 'KBLI codes,', len(ntb_range), 'NTB ranges')
